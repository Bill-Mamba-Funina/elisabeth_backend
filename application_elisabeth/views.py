from decimal import Decimal
from django.utils.dateparse import parse_date

from django.db.models import Count

from django.http import JsonResponse, HttpResponse, FileResponse
from django.db import transaction
from django.db.models import Sum
from django.db.models.functions import TruncMonth
from django.http import JsonResponse, HttpResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone

from .services.notification_service import send_notification

from rest_framework import status, viewsets
from rest_framework.parsers import MultiPartParser, FormParser, JSONParser
from rest_framework.response import Response

from .models import Hall, HallImage, HallVideo

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

from rest_framework import serializers, status, viewsets
from rest_framework.decorators import (
    action,
    api_view,
    permission_classes,
)
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from .models import (
    Client,
    Hall,
    Material,
    Reservation,
    FinancialAccount,
    Payment,
    CashMovement,
    Expense,
    Contract,
    Personnel,
    Refund,
    Tarif,
)

from .serializers import (
    ClientSerializer,
    HallSerializer,
    MaterialSerializer,
    ReservationSerializer,
    FinancialAccountSerializer,
    PaymentSerializer,
    CashMovementSerializer,
    ExpenseSerializer,
    ContractSerializer,
    PersonnelSerializer,
    HallSerializer,
    HallImageSerializer,
    HallVideoSerializer,
    TarifSerializer,
    RefundSerializer,
)

from .services.pdf_service import generate_payment_receipt_pdf

from django.db.models import (
    Count,
    DecimalField,
    F,
    Q,
    Sum,
)


# ============================================================
# CLIENTS
# ============================================================

class ClientViewSet(viewsets.ModelViewSet):

    queryset = Client.objects.all().order_by("full_name")
    serializer_class = ClientSerializer
    permission_classes = [IsAuthenticated]

    def perform_create(self, serializer):

        client = serializer.save()

        send_notification(
            subject="Nouveau client",
            message=(
                f"Un nouveau client a été créé.\n\n"
                f"Nom : {client.full_name}\n"
                f"Téléphone : {client.phone}"
            ),
        )

# ============================================================
# SALLES
# ============================================================


class HallViewSet(viewsets.ModelViewSet):
    queryset = Hall.objects.prefetch_related(
        "images",
        "videos",
    ).all()

    serializer_class = HallSerializer

    parser_classes = [
        MultiPartParser, 
        FormParser,
        JSONParser,
    ]

    permission_classes = [IsAuthenticated]

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(
            data=request.data
        )

        serializer.is_valid(raise_exception=True)

        hall = serializer.save()

        # =====================================================
        # IMAGES
        # =====================================================

        images = request.FILES.getlist("images")

        for image in images:
            HallImage.objects.create(
                hall=hall,
                image=image,
            )

        # =====================================================
        # VIDEOS
        # =====================================================

        videos = request.FILES.getlist("videos")

        for video in videos:
            HallVideo.objects.create(
                hall=hall,
                video=video,
            )

        # =====================================================
        # REPONSE
        # =====================================================

        response_serializer = self.get_serializer(
            hall
        )

        return Response(
            response_serializer.data,
            status=status.HTTP_201_CREATED,
        )

    def update(self, request, *args, **kwargs):
        partial = kwargs.pop(
            "partial",
            False,
        )

        instance = self.get_object()

        serializer = self.get_serializer(
            instance,
            data=request.data,
            partial=partial,
        )

        serializer.is_valid(
            raise_exception=True
        )

        hall = serializer.save()

        # =====================================================
        # NOUVELLES IMAGES
        # =====================================================

        images = request.FILES.getlist(
            "images"
        )

        for image in images:
            HallImage.objects.create(
                hall=hall,
                image=image,
            )

        # =====================================================
        # NOUVELLES VIDEOS
        # =====================================================

        videos = request.FILES.getlist(
            "videos"
        )

        for video in videos:
            HallVideo.objects.create(
                hall=hall,
                video=video,
            )

        response_serializer = self.get_serializer(
            hall
        )

        return Response(
            response_serializer.data
        )



class HallImageViewSet(viewsets.ModelViewSet):
    queryset = HallImage.objects.all().select_related("hall")
    serializer_class = HallImageSerializer

    parser_classes = [
        MultiPartParser,
        FormParser,
        JSONParser,
    ]


class HallVideoViewSet(viewsets.ModelViewSet):
    queryset = HallVideo.objects.all().select_related("hall")
    serializer_class = HallVideoSerializer

    parser_classes = [
        MultiPartParser,
        FormParser,
        JSONParser,
    ]



# ============================================================
# MATERIEL
# ============================================================

class MaterialViewSet(viewsets.ModelViewSet):
    queryset = Material.objects.all().order_by("name")
    serializer_class = MaterialSerializer
    permission_classes = [IsAuthenticated]



# ============================================================
# COMPTES FINANCIERS
# ============================================================

class FinancialAccountViewSet(viewsets.ModelViewSet):
    queryset = FinancialAccount.objects.all()
    serializer_class = FinancialAccountSerializer
    permission_classes = [IsAuthenticated]

    @action(
        detail=True,
        methods=["get"],
        url_path="solde"
    )
    def solde(self, request, pk=None):
        account = self.get_object()

        return Response(
            {
                "id": account.id,
                "name": account.name,
                "account_type": account.account_type,
                "balance": account.balance,
            }
        )



class TarifViewSet(viewsets.ModelViewSet):
    queryset = Tarif.objects.all()
    serializer_class = TarifSerializer
    permission_classes = [IsAuthenticated]




# ============================================================
# PAIEMENTS
# ============================================================


class PaymentViewSet(viewsets.ModelViewSet):
    """
    Gestion complète des paiements.

    Fonctions :
    - GET    /payments/
    - GET    /payments/<id>/
    - POST   /payments/
    - PUT    /payments/<id>/
    - PATCH  /payments/<id>/
    - DELETE /payments/<id>/
    - POST   /payments/<id>/valider/
    - POST   /payments/<id>/annuler/
    - GET    /payments/<id>/recu/
    """

    queryset = (
        Payment.objects
        .select_related(
            "reservation",
            "reservation__client",
            "financial_account",
            "created_by",
        )
        .all()
        .order_by("-payment_date", "-id")
    )

    serializer_class = PaymentSerializer
    permission_classes = [IsAuthenticated]

    # ========================================================
    # UTILITAIRES
    # ========================================================

    def _get_payment_for_update(self, pk):
        """
        Verrouille uniquement la ligne Payment.

        IMPORTANT :
        On ne fait PAS de select_related() ici avec
        select_for_update() afin d'éviter PostgreSQL :

        FOR UPDATE ne peut pas être appliqué sur le côté
        possiblement NULL d'une jointure externe.
        """

        return (
            Payment.objects
            .select_for_update()
            .get(pk=pk)
        )

    def _get_reservation_for_update(self, reservation_id):
        """
        Verrouille séparément la réservation.
        """

        return (
            Reservation.objects
            .select_for_update()
            .get(pk=reservation_id)
        )

    def _get_account_for_update(self, account_id):
        """
        Verrouille séparément le compte financier.
        """

        return (
            FinancialAccount.objects
            .select_for_update()
            .get(pk=account_id)
        )

    def _calculate_paid_amount(self, reservation):
        """
        Calcule le montant total des paiements validés
        pour une réservation.
        """

        result = (
            Payment.objects
            .filter(
                reservation=reservation,
                status=Payment.Status.VALIDE,
            )
            .aggregate(
                total=Sum("amount")
            )
        )

        return (
            result.get("total")
            or Decimal("0.00")
        )

    def _calculate_remaining_amount(self, reservation):
        """
        Calcule le reste à payer.

        Formule :

        total réservation
        - paiements validés
        + remboursements
        """

        total = (
            reservation.total_amount
            or Decimal("0.00")
        )

        already_paid = self._calculate_paid_amount(
            reservation
        )

        refunded = (
            reservation.refunded_amount
            or Decimal("0.00")
        )

        remaining = (
            total
            - already_paid
            + refunded
        )

        if remaining < Decimal("0.00"):
            remaining = Decimal("0.00")

        return remaining

    # ========================================================
    # CREATION
    # ========================================================

    @transaction.atomic
    def perform_create(self, serializer):

        validated_data = serializer.validated_data

        reservation = validated_data.get(
            "reservation"
        )

        idempotency_key = validated_data.get(
            "idempotency_key"
        )

        # ----------------------------------------------------
        # PROTECTION CONTRE LE DOUBLE POST
        # ----------------------------------------------------

        if idempotency_key:

            existing_payment = (
                Payment.objects
                .filter(
                    idempotency_key=idempotency_key
                )
                .first()
            )

            if existing_payment:

                raise serializers.ValidationError({
                    "detail": (
                        "Ce paiement a déjà été enregistré."
                    ),
                    "payment_id": existing_payment.id,
                })

        # ----------------------------------------------------
        # VERROUILLAGE RESERVATION
        # ----------------------------------------------------

        locked_reservation = None

        if reservation:

            locked_reservation = (
                Reservation.objects
                .select_for_update()
                .get(
                    pk=reservation.pk
                )
            )

            total = (
                locked_reservation.total_amount
                or Decimal("0.00")
            )

            already_paid = self._calculate_paid_amount(
                locked_reservation
            )

            refunded = (
                locked_reservation.refunded_amount
                or Decimal("0.00")
            )

            remaining = (
                total
                - already_paid
                + refunded
            )

            if remaining < Decimal("0.00"):
                remaining = Decimal("0.00")

            amount = (
                validated_data.get("amount")
                or Decimal("0.00")
            )

            status_value = validated_data.get(
                "status",
                Payment.Status.EN_ATTENTE,
            )

            # ------------------------------------------------
            # MONTANT POSITIF
            # ------------------------------------------------

            if amount <= Decimal("0.00"):

                raise serializers.ValidationError({
                    "amount": (
                        "Le montant du paiement doit "
                        "être supérieur à zéro."
                    )
                })

            # ------------------------------------------------
            # PAIEMENT ANNULE INTERDIT A LA CREATION
            # ------------------------------------------------

            if status_value == Payment.Status.ANNULE:

                raise serializers.ValidationError({
                    "status": (
                        "Un paiement ne peut pas être "
                        "créé directement avec le statut "
                        "ANNULÉ."
                    )
                })

            # ------------------------------------------------
            # RESERVATION DEJA PAYEE
            # ------------------------------------------------

            if remaining <= Decimal("0.00"):

                raise serializers.ValidationError({
                    "amount": (
                        "Cette réservation est déjà "
                        "entièrement payée."
                    )
                })

            # ------------------------------------------------
            # DEPASSEMENT
            # ------------------------------------------------

            if amount > remaining:

                raise serializers.ValidationError({
                    "amount": (
                        "Le montant maximum autorisé est "
                        f"{remaining} $."
                    )
                })

            # ------------------------------------------------
            # UTILISER LA RESERVATION VERROUILLEE
            # ------------------------------------------------

            validated_data["reservation"] = (
                locked_reservation
            )

        else:

            # ------------------------------------------------
            # PAIEMENT SANS RESERVATION
            # ------------------------------------------------

            amount = (
                validated_data.get("amount")
                or Decimal("0.00")
            )

            if amount <= Decimal("0.00"):

                raise serializers.ValidationError({
                    "amount": (
                        "Le montant du paiement doit "
                        "être supérieur à zéro."
                    )
                })

        # ----------------------------------------------------
        # REFERENCE UNIQUE
        # ----------------------------------------------------

        reference = validated_data.get(
            "reference"
        )

        if reference:

            existing_reference = (
                Payment.objects
                .filter(
                    reference=reference
                )
                .first()
            )

            if existing_reference:

                raise serializers.ValidationError({
                    "reference": (
                        "Cette référence de paiement "
                        "est déjà utilisée."
                    ),
                    "payment_id": existing_reference.id,
                })

        # ----------------------------------------------------
        # CREATION
        # ----------------------------------------------------

        serializer.save(
            created_by=self.request.user
        )

    # ========================================================
    # MODIFICATION
    # ========================================================

    @transaction.atomic
    def perform_update(self, serializer):

        payment = (
            Payment.objects
            .select_for_update()
            .get(
                pk=serializer.instance.pk
            )
        )

        # ----------------------------------------------------
        # PAIEMENT VALIDE
        # ----------------------------------------------------

        if payment.status == Payment.Status.VALIDE:

            raise serializers.ValidationError({
                "detail": (
                    "Un paiement déjà validé ne peut pas "
                    "être modifié directement. "
                    "Utilisez un remboursement ou "
                    "une correction."
                )
            })

        # ----------------------------------------------------
        # PAIEMENT ANNULE
        # ----------------------------------------------------

        if payment.status == Payment.Status.ANNULE:

            raise serializers.ValidationError({
                "detail": (
                    "Un paiement annulé ne peut pas "
                    "être modifié."
                )
            })

        # ----------------------------------------------------
        # NOUVELLE RESERVATION
        # ----------------------------------------------------

        reservation = serializer.validated_data.get(
            "reservation",
            payment.reservation,
        )

        amount = serializer.validated_data.get(
            "amount",
            payment.amount,
        )

        if amount is None or amount <= Decimal("0.00"):

            raise serializers.ValidationError({
                "amount": (
                    "Le montant du paiement doit "
                    "être supérieur à zéro."
                )
            })

        # ----------------------------------------------------
        # VERIFICATION RESERVATION
        # ----------------------------------------------------

        if reservation:

            locked_reservation = (
                Reservation.objects
                .select_for_update()
                .get(
                    pk=reservation.pk
                )
            )

            total = (
                locked_reservation.total_amount
                or Decimal("0.00")
            )

            other_paid = (
                Payment.objects
                .filter(
                    reservation=locked_reservation,
                    status=Payment.Status.VALIDE,
                )
                .exclude(
                    pk=payment.pk
                )
                .aggregate(
                    total=Sum("amount")
                )
                .get("total")
                or Decimal("0.00")
            )

            refunded = (
                locked_reservation.refunded_amount
                or Decimal("0.00")
            )

            remaining = (
                total
                - other_paid
                + refunded
            )

            if remaining < Decimal("0.00"):
                remaining = Decimal("0.00")

            if amount > remaining:

                raise serializers.ValidationError({
                    "amount": (
                        "Le montant maximum autorisé est "
                        f"{remaining} $."
                    )
                })

            serializer.validated_data[
                "reservation"
            ] = locked_reservation

        # ----------------------------------------------------
        # REFERENCE UNIQUE
        # ----------------------------------------------------

        reference = serializer.validated_data.get(
            "reference",
            payment.reference,
        )

        if reference:

            existing_reference = (
                Payment.objects
                .filter(
                    reference=reference
                )
                .exclude(
                    pk=payment.pk
                )
                .first()
            )

            if existing_reference:

                raise serializers.ValidationError({
                    "reference": (
                        "Cette référence de paiement "
                        "est déjà utilisée."
                    )
                })

        serializer.save()

    # ========================================================
    # SUPPRESSION
    # ========================================================

    @transaction.atomic
    def perform_destroy(self, instance):

        payment = (
            Payment.objects
            .select_for_update()
            .get(
                pk=instance.pk
            )
        )

        # ----------------------------------------------------
        # PAIEMENT VALIDE
        # ----------------------------------------------------

        if payment.status == Payment.Status.VALIDE:

            raise serializers.ValidationError({
                "detail": (
                    "Un paiement validé ne peut pas être "
                    "supprimé. Utilisez un remboursement "
                    "ou une correction."
                )
            })

        # ----------------------------------------------------
        # PAIEMENT ANNULE
        # ----------------------------------------------------

        if payment.status == Payment.Status.ANNULE:

            raise serializers.ValidationError({
                "detail": (
                    "Un paiement annulé ne peut pas "
                    "être supprimé."
                )
            })

        payment.delete()

    # ========================================================
    # VALIDER UN PAIEMENT
    # ========================================================

    @action(
        detail=True,
        methods=["post"],
        url_path="valider",
    )
    @transaction.atomic
    def valider(self, request, pk=None):

        # ----------------------------------------------------
        # VERROUILLER UNIQUEMENT LE PAIEMENT
        # ----------------------------------------------------

        payment = self._get_payment_for_update(pk)

        # ----------------------------------------------------
        # DEJA VALIDE
        # ----------------------------------------------------

        if payment.status == Payment.Status.VALIDE:

            payment = (
                Payment.objects
                .select_related(
                    "reservation",
                    "reservation__client",
                    "financial_account",
                    "created_by",
                )
                .get(pk=payment.pk)
            )

            return Response(
                {
                    "detail": (
                        "Ce paiement est déjà validé."
                    ),
                    "payment": PaymentSerializer(
                        payment,
                        context={
                            "request": request,
                        },
                    ).data,
                },
                status=status.HTTP_200_OK,
            )

        # ----------------------------------------------------
        # ANNULE
        # ----------------------------------------------------

        if payment.status == Payment.Status.ANNULE:

            return Response(
                {
                    "detail": (
                        "Un paiement annulé ne peut pas "
                        "être validé."
                    )
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        # ----------------------------------------------------
        # VERIFIER MONTANT
        # ----------------------------------------------------

        if (
            payment.amount is None
            or payment.amount <= Decimal("0.00")
        ):

            return Response(
                {
                    "detail": (
                        "Le montant du paiement doit "
                        "être supérieur à zéro."
                    )
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        # ----------------------------------------------------
        # RESERVATION
        # ----------------------------------------------------

        reservation = None

        if payment.reservation_id:

            reservation = (
                Reservation.objects
                .select_for_update()
                .get(
                    pk=payment.reservation_id
                )
            )

            total = (
                reservation.total_amount
                or Decimal("0.00")
            )

            # ------------------------------------------------
            # AUTRES PAIEMENTS VALIDES
            # ------------------------------------------------

            other_paid = (
                Payment.objects
                .filter(
                    reservation_id=reservation.pk,
                    status=Payment.Status.VALIDE,
                )
                .exclude(
                    pk=payment.pk
                )
                .aggregate(
                    total=Sum("amount")
                )
                .get("total")
                or Decimal("0.00")
            )

            # ------------------------------------------------
            # REMBOURSEMENTS
            # ------------------------------------------------

            refunded = (
                reservation.refunded_amount
                or Decimal("0.00")
            )

            # ------------------------------------------------
            # RESTE DISPONIBLE
            # ------------------------------------------------

            remaining = (
                total
                - other_paid
                + refunded
            )

            if remaining < Decimal("0.00"):
                remaining = Decimal("0.00")

            # ------------------------------------------------
            # RESERVATION DEJA PAYEE
            # ------------------------------------------------

            if remaining <= Decimal("0.00"):

                return Response(
                    {
                        "detail": (
                            "Cette réservation est déjà "
                            "entièrement payée."
                        )
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            # ------------------------------------------------
            # DEPASSEMENT
            # ------------------------------------------------

            if payment.amount > remaining:

                return Response(
                    {
                        "detail": (
                            "Le montant du paiement "
                            f"({payment.amount} $) dépasse "
                            "le reste à payer "
                            f"({remaining} $)."
                        )
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

        # ====================================================
        # VALIDATION DU PAIEMENT
        # ====================================================

        payment.status = Payment.Status.VALIDE

        payment.save(
            update_fields=[
                "status",
            ]
        )

        # ====================================================
        # MOUVEMENT FINANCIER
        # ====================================================

        if payment.financial_account_id:

            account = self._get_account_for_update(
                payment.financial_account_id
            )

            # ------------------------------------------------
            # EVITER LE DOUBLE MOUVEMENT
            # ------------------------------------------------

            mouvement_existe = (
                CashMovement.objects
                .filter(
                    payment_id=payment.pk,
                    movement_type=(
                        CashMovement.MovementType.ENTREE
                    ),
                )
                .exists()
            )

            if not mouvement_existe:

                account.balance += payment.amount

                account.save(
                    update_fields=[
                        "balance",
                    ]
                )

                CashMovement.objects.create(
                    account=account,
                    payment=payment,
                    reservation=payment.reservation,
                    movement_type=(
                        CashMovement.MovementType.ENTREE
                    ),
                    amount=payment.amount,
                    description=(
                        f"Paiement validé "
                        f"#{payment.id}"
                    ),
                    created_by=request.user,
                )

        # ====================================================
        # RECALCUL FINANCIER RESERVATION
        # ====================================================

        if reservation:

            reservation.recalculate_financials()

        # ====================================================
        # GENERATION DU RECU PDF
        # ====================================================

        try:

            generate_payment_receipt_pdf(
                payment
            )

        except Exception as exc:

            print(
                "ERREUR GENERATION RECU PDF :",
                exc,
            )

        # ====================================================
        # RECHARGER POUR SERIALIZER
        # ====================================================

        payment = (
            Payment.objects
            .select_related(
                "reservation",
                "reservation__client",
                "financial_account",
                "created_by",
            )
            .get(
                pk=payment.pk
            )
        )

        # ====================================================
        # REPONSE
        # ====================================================

        return Response(
            PaymentSerializer(
                payment,
                context={
                    "request": request,
                },
            ).data,
            status=status.HTTP_200_OK,
        )

    # ========================================================
    # ANNULER UN PAIEMENT
    # ========================================================

    @action(
        detail=True,
        methods=["post"],
        url_path="annuler",
    )
    @transaction.atomic
    def annuler(self, request, pk=None):

        payment = self._get_payment_for_update(pk)

        # ----------------------------------------------------
        # DEJA VALIDE
        # ----------------------------------------------------

        if payment.status == Payment.Status.VALIDE:

            return Response(
                {
                    "detail": (
                        "Un paiement déjà validé ne peut "
                        "pas être annulé directement. "
                        "Utilisez une procédure de "
                        "remboursement ou de correction."
                    )
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        # ----------------------------------------------------
        # DEJA ANNULE
        # ----------------------------------------------------

        if payment.status == Payment.Status.ANNULE:

            return Response(
                {
                    "detail": (
                        "Ce paiement est déjà annulé."
                    )
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        # ----------------------------------------------------
        # ANNULATION
        # ----------------------------------------------------

        payment.status = Payment.Status.ANNULE

        payment.save(
            update_fields=[
                "status",
            ]
        )

        # ----------------------------------------------------
        # REPONSE
        # ----------------------------------------------------

        payment = (
            Payment.objects
            .select_related(
                "reservation",
                "reservation__client",
                "financial_account",
                "created_by",
            )
            .get(
                pk=payment.pk
            )
        )

        return Response(
            PaymentSerializer(
                payment,
                context={
                    "request": request,
                },
            ).data,
            status=status.HTTP_200_OK,
        )

    # ========================================================
    # AFFICHER / TELECHARGER LE RECU PDF
    # ========================================================

    @action(
        detail=True,
        methods=["get"],
        url_path="recu",
    )
    def recu(self, request, pk=None):

        # ----------------------------------------------------
        # RECUPERER LE PAIEMENT
        # ----------------------------------------------------

        payment = (
            Payment.objects
            .select_related(
                "reservation",
                "reservation__client",
                "financial_account",
                "created_by",
            )
            .get(
                pk=pk
            )
        )

        # ----------------------------------------------------
        # VERIFICATION STATUT
        # ----------------------------------------------------

        if payment.status != Payment.Status.VALIDE:

            return Response(
                {
                    "detail": (
                        "Le reçu ne peut être généré "
                        "que pour un paiement validé."
                    )
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        # ----------------------------------------------------
        # GENERATION PDF
        # ----------------------------------------------------

        try:

            pdf_content = (
                generate_payment_receipt_pdf(
                    payment
                )
            )

        except Exception as exc:

            print(
                "ERREUR GENERATION RECU PDF :",
                exc,
            )

            return Response(
                {
                    "detail": (
                        "Impossible de générer "
                        "le reçu PDF."
                    ),
                    "error": str(exc),
                },
                status=status.HTTP_500_INTERNAL_SERVER_ERROR,
            )

        # ----------------------------------------------------
        # REPONSE HTTP PDF
        # ----------------------------------------------------

        response = HttpResponse(
            pdf_content,
            content_type="application/pdf",
        )

        response[
            "Content-Disposition"
        ] = (
            f'inline; filename="recu-paiement-'
            f'{payment.id}.pdf"'
        )

        response[
            "Content-Length"
        ] = str(
            len(pdf_content)
        )

        return response

# ============================================================
# RESERVATIONS
# ============================================================

class ReservationViewSet(viewsets.ModelViewSet):

    queryset = (
        Reservation.objects
        .select_related(
            "client",
            "hall",
            "tarif",
        )
        .prefetch_related(
            "payments",
            "refunds",
        )
        .order_by("-event_date")
    )

    serializer_class = ReservationSerializer
    permission_classes = [IsAuthenticated]

    # ========================================================
    # CREATION
    # ========================================================

    def perform_create(self, serializer):

        reservation = serializer.save(
            created_by=self.request.user
        )

        client_name = (
            reservation.client.full_name
            if reservation.client
            else "Non renseigné"
        )

        hall_name = (
            reservation.hall.name
            if reservation.hall
            else "Non renseignée"
        )

        send_notification(
            subject="Nouvelle réservation",
            message=(
                "Une nouvelle réservation a été créée.\n\n"
                f"Numéro : {reservation.reservation_number}\n"
                f"Client : {client_name}\n"
                f"Salle : {hall_name}\n"
                f"Événement : "
                f"{reservation.event_type or 'Non renseigné'}\n"
                f"Date : {reservation.event_date}\n"
                f"Statut : "
                f"{reservation.get_status_display()}\n"
                f"Montant : "
                f"{reservation.total_amount or Decimal('0.00')} $"
            ),
        )

    # ========================================================
    # MODIFICATION
    # ========================================================

    def perform_update(self, serializer):

        old_status = serializer.instance.status

        reservation = serializer.save()

        # ----------------------------------------------------
        # RESERVATION CONFIRMEE
        # ----------------------------------------------------

        if (
            old_status != Reservation.Status.CONFIRMEE
            and reservation.status
            == Reservation.Status.CONFIRMEE
        ):

            client_name = (
                reservation.client.full_name
                if reservation.client
                else "Non renseigné"
            )

            hall_name = (
                reservation.hall.name
                if reservation.hall
                else "Non renseignée"
            )

            send_notification(
                subject="Réservation confirmée",
                message=(
                    f"La réservation "
                    f"{reservation.reservation_number} "
                    f"a été confirmée.\n\n"
                    f"Client : {client_name}\n"
                    f"Salle : {hall_name}\n"
                    f"Date : {reservation.event_date}"
                ),
            )

    # ========================================================
    # CONFIRMER
    # ========================================================

    @action(
        detail=True,
        methods=["post"],
        url_path="confirmer",
    )
    def confirmer(self, request, pk=None):

        reservation = self.get_object()

        if (
            reservation.status
            == Reservation.Status.ANNULEE
        ):
            return Response(
                {
                    "detail": (
                        "Une réservation annulée "
                        "ne peut pas être confirmée."
                    )
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        reservation.status = Reservation.Status.CONFIRMEE

        reservation.save()

        client_name = (
            reservation.client.full_name
            if reservation.client
            else "Non renseigné"
        )

        hall_name = (
            reservation.hall.name
            if reservation.hall
            else "Non renseignée"
        )

        send_notification(
            subject="Réservation confirmée",
            message=(
                f"La réservation "
                f"{reservation.reservation_number} "
                f"est maintenant confirmée.\n\n"
                f"Client : {client_name}\n"
                f"Salle : {hall_name}\n"
                f"Date : {reservation.event_date}"
            ),
        )

        return Response(
            ReservationSerializer(
                reservation,
                context={
                    "request": request,
                },
            ).data,
            status=status.HTTP_200_OK,
        )

    # ========================================================
    # ANNULER
    # ========================================================

    @action(
        detail=True,
        methods=["post"],
        url_path="annuler",
    )
    def annuler(self, request, pk=None):

        reservation = self.get_object()

        if (
            reservation.status
            == Reservation.Status.ANNULEE
        ):
            return Response(
                {
                    "detail": (
                        "Cette réservation est déjà annulée."
                    )
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        reservation.status = Reservation.Status.ANNULEE

        reservation.save()

        client_name = (
            reservation.client.full_name
            if reservation.client
            else "Non renseigné"
        )

        hall_name = (
            reservation.hall.name
            if reservation.hall
            else "Non renseignée"
        )

        send_notification(
            subject="Réservation annulée",
            message=(
                f"La réservation "
                f"{reservation.reservation_number} "
                f"a été annulée.\n\n"
                f"Client : {client_name}\n"
                f"Salle : {hall_name}\n"
                f"Date : {reservation.event_date}"
            ),
        )

        return Response(
            ReservationSerializer(
                reservation,
                context={
                    "request": request,
                },
            ).data,
            status=status.HTTP_200_OK,
        )



# ============================================================
# DEPENSES
# ============================================================

class ExpenseViewSet(viewsets.ModelViewSet):

    queryset = (
        Expense.objects
        .select_related("created_by")
        .all()
        .order_by("-expense_date", "-id")
    )

    serializer_class = ExpenseSerializer
    permission_classes = [IsAuthenticated]

    # ========================================================
    # CREATION
    # ========================================================

    @transaction.atomic
    def perform_create(self, serializer):

        serializer.save(
            created_by=self.request.user
        )

        expense = serializer.instance

        send_notification(
            subject="Nouvelle dépense",
            message=(
                f"Une nouvelle dépense a été enregistrée.\n\n"
                f"Titre : {expense.title}\n"
                f"Catégorie : "
                f"{expense.get_category_display()}\n"
                f"Montant : {expense.amount} $\n"
                f"Statut : "
                f"{expense.get_status_display()}"
            ),
        )

    # ========================================================
    # MODIFICATION
    # ========================================================

    @transaction.atomic
    def perform_update(self, serializer):

        serializer.save()

    # ========================================================
    # PAYER UNE DEPENSE
    # ========================================================

    @action(
        detail=True,
        methods=["post"],
        url_path="payer",
    )
    @transaction.atomic
    def payer(self, request, pk=None):

        expense = (
            Expense.objects
            .select_for_update()
            .select_related("created_by")
            .get(pk=pk)
        )

        # ----------------------------------------------------
        # DEJA PAYEE
        # ----------------------------------------------------

        if expense.status == Expense.Status.PAYEE:

            return Response(
                {
                    "detail": (
                        "Cette dépense est déjà payée."
                    )
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        # ----------------------------------------------------
        # ANNULEE
        # ----------------------------------------------------

        if expense.status == Expense.Status.ANNULEE:

            return Response(
                {
                    "detail": (
                        "Une dépense annulée "
                        "ne peut pas être payée."
                    )
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        # ----------------------------------------------------
        # PASSAGE A PAYEE
        # ----------------------------------------------------
        #
        # Le modèle Expense.save() s'occupe automatiquement :
        #
        # - de trouver la caisse active ;
        # - vérifier le solde ;
        # - diminuer la caisse ;
        # - créer CashMovement SORTIE.
        #

        expense.status = Expense.Status.PAYEE

        expense.save(
            update_fields=[
                "status",
            ]
        )

        send_notification(
            subject="Dépense payée",
            message=(
                f"La dépense #{expense.id} a été payée.\n\n"
                f"Titre : {expense.title}\n"
                f"Montant : {expense.amount} $"
            ),
        )

        return Response(
            ExpenseSerializer(
                expense,
                context={
                    "request": request,
                },
            ).data,
            status=status.HTTP_200_OK,
        )

# ============================================================
# MOUVEMENTS FINANCIERS
# ============================================================

class CashMovementViewSet(viewsets.ReadOnlyModelViewSet):
    queryset = CashMovement.objects.select_related(
        "account",
        "payment",
        "reservation",
        "created_by",
    ).all()

    serializer_class = CashMovementSerializer
    permission_classes = [IsAuthenticated]


# ============================================================
# CONTRATS
# ============================================================

class ContractViewSet(viewsets.ModelViewSet):
    queryset = Contract.objects.select_related(
        "reservation",
        "reservation__client",
    ).all()

    serializer_class = ContractSerializer
    permission_classes = [IsAuthenticated]

    def perform_create(self, serializer):
        serializer.save(
            uploaded_by=self.request.user
        )



# ============================================================
# CALENDRIER
# ============================================================

@api_view(["GET"])
@permission_classes([IsAuthenticated])
def calendar_view(request, year, month):
    """
    Retourne les réservations du mois demandé
    pour alimenter le calendrier frontend.
    """

    reservations = (
        Reservation.objects
        .filter(
            event_date__year=year,
            event_date__month=month,
        )
        .select_related("client", "hall")
        .order_by("event_date", "start_time")
    )

    results = []

    for reservation in reservations:
        results.append({
            "id": reservation.id,
            "reservation_number": reservation.reservation_number,

            "client": (
                str(reservation.client)
                if reservation.client
                else None
            ),

            "hall": (
                str(reservation.hall)
                if reservation.hall
                else None
            ),

            "event_type": reservation.event_type or "",

            "date": (
                reservation.event_date.isoformat()
                if reservation.event_date
                else None
            ),

            "start_time": (
                reservation.start_time.strftime("%H:%M")
                if reservation.start_time
                else ""
            ),

            "end_time": (
                reservation.end_time.strftime("%H:%M")
                if reservation.end_time
                else ""
            ),

            "status": reservation.status,

            "payment_status": getattr(
                reservation,
                "payment_status",
                ""
            ),
        })

    return JsonResponse({
        "year": year,
        "month": month,
        "count": len(results),
        "results": results,
    })




class PersonnelViewSet(viewsets.ModelViewSet):

    queryset = Personnel.objects.all().order_by(
        "nom",
        "prenom",
    )

    serializer_class = PersonnelSerializer
    permission_classes = [IsAuthenticated]

    def perform_create(self, serializer):

        personnel = serializer.save()

        send_notification(
            subject="Nouveau personnel",
            message=(
                f"Un nouveau membre du personnel "
                f"a été ajouté.\n\n"
                f"Nom : "
                f"{personnel.nom} "
                f"{personnel.prenom}"
            ),
        )



# ============================================================
# TABLEAU DE BORD / RAPPORTS
# ============================================================






def _dashboard_filters(request):
    """
    Récupère les filtres envoyés par le frontend.

    Filtres acceptés :

        date=2026-10-01
        jour=1
        mois=10
        annee=2026

        date_debut=2026-10-01
        date_fin=2026-10-31
    """

    params = request.query_params

    date_value = parse_date(
        (params.get("date") or "").strip()
    )

    date_debut = parse_date(
        (params.get("date_debut") or "").strip()
    )

    date_fin = parse_date(
        (params.get("date_fin") or "").strip()
    )

    try:
        jour = int(params.get("jour"))
        if jour < 1 or jour > 31:
            jour = None
    except (TypeError, ValueError):
        jour = None

    try:
        mois = int(params.get("mois"))
        if mois < 1 or mois > 12:
            mois = None
    except (TypeError, ValueError):
        mois = None

    try:
        annee = int(params.get("annee"))
        if annee < 2000 or annee > 2100:
            annee = None
    except (TypeError, ValueError):
        annee = None

    return {
        "date": date_value,
        "jour": jour,
        "mois": mois,
        "annee": annee,
        "date_debut": date_debut,
        "date_fin": date_fin,
    }


def _filter_dashboard_querysets(request):
    """
    Applique les mêmes filtres à toutes les sources
    du dashboard.
    """

    filters_data = _dashboard_filters(request)

    reservations = Reservation.objects.all()

    payments = Payment.objects.filter(
        status=Payment.Status.VALIDE
    )

    expenses = Expense.objects.all()

    refunds = Refund.objects.all()

    movements = CashMovement.objects.all()

    # ========================================================
    # DATE EXACTE
    # ========================================================

    if filters_data["date"]:

        reservations = reservations.filter(
            event_date=filters_data["date"]
        )

        payments = payments.filter(
            payment_date__date=filters_data["date"]
        )

        expenses = expenses.filter(
            expense_date__date=filters_data["date"]
        )

        refunds = refunds.filter(
            refund_date__date=filters_data["date"]
        )

        movements = movements.filter(
            created_at__date=filters_data["date"]
        )

    # ========================================================
    # JOUR
    # ========================================================

    if filters_data["jour"]:

        reservations = reservations.filter(
            event_date__day=filters_data["jour"]
        )

        payments = payments.filter(
            payment_date__day=filters_data["jour"]
        )

        expenses = expenses.filter(
            expense_date__day=filters_data["jour"]
        )

        refunds = refunds.filter(
            refund_date__day=filters_data["jour"]
        )

        movements = movements.filter(
            created_at__day=filters_data["jour"]
        )

    # ========================================================
    # MOIS
    # ========================================================

    if filters_data["mois"]:

        reservations = reservations.filter(
            event_date__month=filters_data["mois"]
        )

        payments = payments.filter(
            payment_date__month=filters_data["mois"]
        )

        expenses = expenses.filter(
            expense_date__month=filters_data["mois"]
        )

        refunds = refunds.filter(
            refund_date__month=filters_data["mois"]
        )

        movements = movements.filter(
            created_at__month=filters_data["mois"]
        )

    # ========================================================
    # ANNEE
    # ========================================================

    if filters_data["annee"]:

        reservations = reservations.filter(
            event_date__year=filters_data["annee"]
        )

        payments = payments.filter(
            payment_date__year=filters_data["annee"]
        )

        expenses = expenses.filter(
            expense_date__year=filters_data["annee"]
        )

        refunds = refunds.filter(
            refund_date__year=filters_data["annee"]
        )

        movements = movements.filter(
            created_at__year=filters_data["annee"]
        )

    # ========================================================
    # DATE DEBUT
    # ========================================================

    if filters_data["date_debut"]:

        reservations = reservations.filter(
            event_date__gte=filters_data["date_debut"]
        )

        payments = payments.filter(
            payment_date__date__gte=filters_data["date_debut"]
        )

        expenses = expenses.filter(
            expense_date__date__gte=filters_data["date_debut"]
        )

        refunds = refunds.filter(
            refund_date__date__gte=filters_data["date_debut"]
        )

        movements = movements.filter(
            created_at__date__gte=filters_data["date_debut"]
        )

    # ========================================================
    # DATE FIN
    # ========================================================

    if filters_data["date_fin"]:

        reservations = reservations.filter(
            event_date__lte=filters_data["date_fin"]
        )

        payments = payments.filter(
            payment_date__date__lte=filters_data["date_fin"]
        )

        expenses = expenses.filter(
            expense_date__date__lte=filters_data["date_fin"]
        )

        refunds = refunds.filter(
            refund_date__date__lte=filters_data["date_fin"]
        )

        movements = movements.filter(
            created_at__date__lte=filters_data["date_fin"]
        )

    return (
        reservations,
        payments,
        expenses,
        refunds,
        movements,
        filters_data,
    )

def _dashboard_payload(request):
    """
    Construit toutes les données du dashboard
    directement depuis PostgreSQL.
    """

    (
        reservations,
        payments,
        expenses,
        refunds,
        movements,
        filters_data,
    ) = _filter_dashboard_querysets(request)

    reservations = reservations.select_related(
        "client",
        "hall",
        "tarif",
    )

    payments = payments.select_related(
        "reservation",
        "reservation__client",
        "financial_account",
    )

    expenses = expenses.select_related(
        "created_by",
    )

    refunds = refunds.select_related(
        "payment",
        "reservation",
        "financial_account",
    )

    movements = movements.select_related(
        "account",
        "payment",
        "reservation",
        "created_by",
    )

    # ========================================================
    # RESERVATIONS
    # ========================================================

    total_reservations = reservations.count()

    reservations_annulees = reservations.filter(
        status=Reservation.Status.ANNULEE
    ).count()

    reservations_actives = reservations.exclude(
        status=Reservation.Status.ANNULEE
    ).count()

    reservations_terminees = reservations.filter(
        status__in=[
            Reservation.Status.TERMINEE,
            Reservation.Status.CLOTUREE,
        ]
    ).count()

    # ========================================================
    # CHIFFRE D'AFFAIRES
    # ========================================================

    chiffre_affaires = Decimal("0.00")

    for reservation in reservations:

        if reservation.status == Reservation.Status.ANNULEE:
            continue

        if reservation.tarif:

            chiffre_affaires += (
                reservation.tarif.amount
                or Decimal("0.00")
            )

    # ========================================================
    # ENCAISSEMENTS
    # ========================================================

    total_encaisse = (
        payments.aggregate(
            total=Sum("amount")
        ).get("total")
        or Decimal("0.00")
    )

    # ========================================================
    # DEPENSES
    # ========================================================

    total_depenses = (
        expenses.aggregate(
            total=Sum("amount")
        ).get("total")
        or Decimal("0.00")
    )

    # ========================================================
    # REMBOURSEMENTS
    # ========================================================

    total_rembourse = (
        refunds.aggregate(
            total=Sum("amount")
        ).get("total")
        or Decimal("0.00")
    )

    # ========================================================
    # RESTE A RECOUVRER
    # ========================================================

    reste_a_recouvrer = Decimal("0.00")

    for reservation in reservations:

        if reservation.status == Reservation.Status.ANNULEE:
            continue

        reste_a_recouvrer += (
            reservation.remaining_amount
            or Decimal("0.00")
        )

    # ========================================================
    # COMPTES FINANCIERS
    # ========================================================

    accounts = list(
        FinancialAccount.objects.filter(
            is_active=True
        ).order_by("name")
    )

    solde_comptes = sum(
        (
            account.balance
            or Decimal("0.00")
        )
        for account in accounts
    )

    comptes = []

    for account in accounts:

        comptes.append({
            "id": account.id,
            "name": account.name,
            "account_type": account.account_type,
            "account_type_label": (
                account.get_account_type_display()
            ),
            "balance": float(
                account.balance
                or Decimal("0.00")
            ),
        })

    # ========================================================
    # MOUVEMENTS
    # ========================================================

    total_entrees = (
        movements.filter(
            movement_type=(
                CashMovement.MovementType.ENTREE
            )
        )
        .aggregate(total=Sum("amount"))
        .get("total")
        or Decimal("0.00")
    )

    total_sorties = (
        movements.filter(
            movement_type__in=[
                CashMovement.MovementType.SORTIE,
            ]
        )
        .aggregate(total=Sum("amount"))
        .get("total")
        or Decimal("0.00")
    )

    # ========================================================
    # RESULTAT NET
    # ========================================================

    resultat_net = (
        total_encaisse
        - total_depenses
        - total_rembourse
    )

    # ========================================================
    # PAIEMENTS PAR MODE
    # ========================================================

    payments_by_method = []

    methods = (
        payments
        .values("method")
        .annotate(
            amount=Sum("amount")
        )
        .order_by("method")
    )

    payment_display_map = dict(
        Payment.Method.choices
    )

    for item in methods:

        method = item["method"]

        payments_by_method.append({
            "method": method,
            "label": payment_display_map.get(
                method,
                method,
            ),
            "amount": float(
                item["amount"]
                or Decimal("0.00")
            ),
        })

    # ========================================================
    # DEPENSES PAR CATEGORIE
    # ========================================================

    expenses_by_category = []

    categories = (
        expenses
        .values("category")
        .annotate(
            amount=Sum("amount")
        )
        .order_by("category")
    )

    expense_category_map = dict(
        getattr(
            Expense,
            "ExpenseType",
            []
        ).choices
        if hasattr(Expense, "ExpenseType")
        else []
    )

    for item in categories:

        category = item["category"]

        expenses_by_category.append({
            "category": category,
            "label": expense_category_map.get(
                category,
                category,
            ),
            "amount": float(
                item["amount"]
                or Decimal("0.00")
            ),
        })

    # ========================================================
    # RESERVATIONS PAR STATUT
    # ========================================================

    reservations_by_status = []

    statuses = (
        reservations
        .values("status")
        .annotate(
            count=Count("id")
        )
        .order_by("status")
    )

    reservation_status_map = dict(
        Reservation.Status.choices
    )

    for item in statuses:

        status_value = item["status"]

        reservations_by_status.append({
            "status": status_value,
            "label": reservation_status_map.get(
                status_value,
                status_value,
            ),
            "count": item["count"],
        })

    # ========================================================
    # RECETTES MENSUELLES
    # ========================================================

    revenues = (
        payments
        .annotate(
            month=TruncMonth("payment_date")
        )
        .values("month")
        .annotate(
            amount=Sum("amount")
        )
        .order_by("month")
    )

    revenus_chart = []

    for item in revenues:

        month = item["month"]

        revenus_chart.append({
            "month": (
                month.strftime("%Y-%m")
                if month
                else ""
            ),
            "amount": float(
                item["amount"]
                or Decimal("0.00")
            ),
        })

    # ========================================================
    # DEPENSES MENSUELLES
    # ========================================================

    expense_months = (
        expenses
        .annotate(
            month=TruncMonth("expense_date")
        )
        .values("month")
        .annotate(
            amount=Sum("amount")
        )
        .order_by("month")
    )

    depenses_chart = []

    for item in expense_months:

        month = item["month"]

        depenses_chart.append({
            "month": (
                month.strftime("%Y-%m")
                if month
                else ""
            ),
            "amount": float(
                item["amount"]
                or Decimal("0.00")
            ),
        })

    # ========================================================
    # REPONSE
    # ========================================================

    return {
        "filters": {
            "date": (
                filters_data["date"].isoformat()
                if filters_data["date"]
                else None
            ),
            "jour": filters_data["jour"],
            "mois": filters_data["mois"],
            "annee": filters_data["annee"],
            "date_debut": (
                filters_data["date_debut"].isoformat()
                if filters_data["date_debut"]
                else None
            ),
            "date_fin": (
                filters_data["date_fin"].isoformat()
                if filters_data["date_fin"]
                else None
            ),
        },

        "summary": {
            "total_reservations": total_reservations,
            "reservations_actives": reservations_actives,
            "reservations_annulees": reservations_annulees,
            "reservations_terminees": reservations_terminees,

            "chiffre_affaires": float(
                chiffre_affaires
            ),

            "total_encaisse": float(
                total_encaisse
            ),

            "total_depenses": float(
                total_depenses
            ),

            "total_rembourse": float(
                total_rembourse
            ),

            "reste_a_recouvrer": float(
                reste_a_recouvrer
            ),

            "solde_comptes": float(
                solde_comptes
            ),

            "resultat_net": float(
                resultat_net
            ),

            "total_entrees": float(
                total_entrees
            ),

            "total_sorties": float(
                total_sorties
            ),
        },

        "accounts": comptes,

        "payments_by_method": (
            payments_by_method
        ),

        "expenses_by_category": (
            expenses_by_category
        ),

        "reservations_by_status": (
            reservations_by_status
        ),

        "monthly": {
            "revenues": revenus_chart,
            "expenses": depenses_chart,
        },
    }


# ============================================================
# API DASHBOARD
# ============================================================

@api_view(["GET"])
@permission_classes([IsAuthenticated])
def dashboard_report(request):

    return Response(
        _dashboard_payload(request),
        status=200,
    )


## ============================================================
# EXPORT EXCEL
# ============================================================

@api_view(["GET"])
@permission_classes([IsAuthenticated])
def dashboard_excel(request):

    (
        reservations,
        payments,
        expenses,
        refunds,
        movements,
        filters_data,
    ) = _filter_dashboard_querysets(request)

    # ========================================================
    # RELATIONS OPTIMISÉES
    # ========================================================

    reservations = reservations.select_related(
        "client",
        "hall",
        "tarif",
    )

    payments = payments.select_related(
        "reservation",
        "reservation__client",
        "financial_account",
    )

    expenses = expenses.select_related(
        "created_by",
    )

    refunds = refunds.select_related(
        "payment",
        "reservation",
        "financial_account",
    )

    movements = movements.select_related(
        "account",
        "reservation",
        "payment",
    )

    # ========================================================
    # CLASSEUR
    # ========================================================

    workbook = Workbook()

    # ========================================================
    # SYNTHÈSE
    # ========================================================

    summary_sheet = workbook.active
    summary_sheet.title = "Synthèse"

    summary_sheet.append([
        "RAPPORT FINANCIER - LA CASA DA FESTA ELISABETH"
    ])

    summary_sheet.append([])

    payload = _dashboard_payload(request)
    summary = payload["summary"]

    summary_rows = [
        (
            "Réservations",
            summary.get("total_reservations", 0),
        ),
        (
            "Réservations actives",
            summary.get("reservations_actives", 0),
        ),
        (
            "Réservations annulées",
            summary.get("reservations_annulees", 0),
        ),
        (
            "Réservations terminées",
            summary.get("reservations_terminees", 0),
        ),
        (
            "Chiffre d'affaires",
            summary.get("chiffre_affaires", 0),
        ),
        (
            "Total encaissé",
            summary.get("total_encaisse", 0),
        ),
        (
            "Total dépenses",
            summary.get("total_depenses", 0),
        ),
        (
            "Total remboursé",
            summary.get("total_rembourse", 0),
        ),
        (
            "Reste à recouvrer",
            summary.get("reste_a_recouvrer", 0),
        ),
        (
            "Solde comptes",
            summary.get("solde_comptes", 0),
        ),
        (
            "Résultat net",
            summary.get("resultat_net", 0),
        ),
        (
            "Total entrées",
            summary.get("total_entrees", 0),
        ),
        (
            "Total sorties",
            summary.get("total_sorties", 0),
        ),
    ]

    summary_sheet.append([
        "Indicateur",
        "Valeur",
    ])

    for label, value in summary_rows:

        summary_sheet.append([
            label,
            value,
        ])

    # ========================================================
    # RÉSERVATIONS
    # ========================================================

    reservation_sheet = workbook.create_sheet(
        "Réservations"
    )

    reservation_sheet.append([
        "N°",
        "Client",
        "Téléphone",
        "Salle",
        "Type",
        "Date",
        "Statut",
        "Paiement",
        "Montant",
        "Payé",
        "Reste",
    ])

    for reservation in reservations:

        reservation_sheet.append([
            reservation.reservation_number,

            (
                reservation.client.full_name
                if reservation.client
                else ""
            ),

            (
                reservation.client.phone
                if reservation.client
                else ""
            ),

            (
                reservation.hall.name
                if reservation.hall
                else ""
            ),

            reservation.event_type or "",

            (
                reservation.event_date.isoformat()
                if reservation.event_date
                else ""
            ),

            reservation.get_status_display(),

            reservation.get_payment_status_display(),

            float(
                reservation.total_amount
                or Decimal("0.00")
            ),

            float(
                reservation.net_paid_amount
                or Decimal("0.00")
            ),

            float(
                reservation.remaining_amount
                or Decimal("0.00")
            ),
        ])

    # ========================================================
    # PAIEMENTS
    # ========================================================

    payment_sheet = workbook.create_sheet(
        "Paiements"
    )

    payment_sheet.append([
        "ID",
        "Réservation",
        "Client",
        "Montant",
        "Date",
        "Mode",
        "Compte",
        "Référence",
        "Statut",
    ])

    for payment in payments:

        payment_sheet.append([
            payment.id,

            (
                payment.reservation.reservation_number
                if payment.reservation
                else ""
            ),

            (
                payment.reservation.client.full_name
                if (
                    payment.reservation
                    and payment.reservation.client
                )
                else ""
            ),

            float(
                payment.amount
                or Decimal("0.00")
            ),

            (
                payment.payment_date.isoformat()
                if payment.payment_date
                else ""
            ),

            (
                payment.get_method_display()
                if hasattr(
                    payment,
                    "get_method_display"
                )
                else payment.method
            ),

            (
                payment.financial_account.name
                if payment.financial_account
                else ""
            ),

            payment.reference or "",

            (
                payment.get_status_display()
                if hasattr(
                    payment,
                    "get_status_display"
                )
                else payment.status
            ),
        ])

    # ========================================================
    # DÉPENSES
    # ========================================================

    expense_sheet = workbook.create_sheet(
        "Dépenses"
    )

    expense_sheet.append([
        "ID",
        "Catégorie",
        "Titre",
        "Notes",
        "Montant",
        "Date",
    ])

    for expense in expenses:

        expense_sheet.append([
            expense.id,

            (
                expense.get_category_display()
                if hasattr(
                    expense,
                    "get_category_display"
                )
                else getattr(
                    expense,
                    "category",
                    "",
                )
            ),

            getattr(
                expense,
                "title",
                "",
            ) or "",

            getattr(
                expense,
                "notes",
                "",
            ) or "",

            float(
                getattr(
                    expense,
                    "amount",
                    Decimal("0.00"),
                )
                or Decimal("0.00")
            ),

            (
                expense.expense_date.isoformat()
                if getattr(
                    expense,
                    "expense_date",
                    None,
                )
                else ""
            ),
        ])

    # ========================================================
    # REMBOURSEMENTS
    # ========================================================

    refund_sheet = workbook.create_sheet(
        "Remboursements"
    )

    refund_sheet.append([
        "ID",
        "Réservation",
        "Paiement",
        "Montant",
        "Date",
        "Mode",
        "Compte",
        "Motif",
        "Référence",
        "Statut",
    ])

    for refund in refunds:

        refund_sheet.append([
            refund.id,

            (
                refund.reservation.reservation_number
                if refund.reservation
                else ""
            ),

            (
                refund.payment.id
                if refund.payment
                else ""
            ),

            float(
                refund.amount
                or Decimal("0.00")
            ),

            (
                refund.refund_date.isoformat()
                if refund.refund_date
                else ""
            ),

            (
                refund.get_method_display()
                if hasattr(
                    refund,
                    "get_method_display"
                )
                else getattr(
                    refund,
                    "method",
                    "",
                )
            ),

            (
                refund.financial_account.name
                if refund.financial_account
                else ""
            ),

            refund.reason or "",

            refund.reference or "",

            (
                refund.get_status_display()
                if hasattr(
                    refund,
                    "get_status_display"
                )
                else getattr(
                    refund,
                    "status",
                    "",
                )
            ),
        ])

    # ========================================================
    # MOUVEMENTS
    # ========================================================

    movement_sheet = workbook.create_sheet(
        "Mouvements"
    )

    movement_sheet.append([
        "ID",
        "Compte",
        "Type",
        "Montant",
        "Description",
        "Date",
        "Réservation",
        "Paiement",
    ])

    for movement in movements:

        movement_sheet.append([
            movement.id,

            (
                movement.account.name
                if movement.account
                else ""
            ),

            (
                movement.get_movement_type_display()
                if hasattr(
                    movement,
                    "get_movement_type_display"
                )
                else getattr(
                    movement,
                    "movement_type",
                    "",
                )
            ),

            float(
                movement.amount
                or Decimal("0.00")
            ),

            getattr(
                movement,
                "description",
                "",
            ) or "",

            (
                movement.created_at.isoformat()
                if movement.created_at
                else ""
            ),

            (
                movement.reservation.reservation_number
                if movement.reservation
                else ""
            ),

            (
                movement.payment.id
                if movement.payment
                else ""
            ),
        ])

    # ========================================================
    # STYLE DES FEUILLES
    # ========================================================

    for worksheet in workbook.worksheets:

        # ----------------------------------------------------
        # En-tête
        # ----------------------------------------------------

        for cell in worksheet[1]:

            cell.font = Font(
                bold=True
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        # ----------------------------------------------------
        # Deuxième ligne pour la synthèse
        # ----------------------------------------------------

        if worksheet.title == "Synthèse":

            for cell in worksheet[3]:

                cell.font = Font(
                    bold=True
                )

                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                )

        # ----------------------------------------------------
        # Largeur automatique des colonnes
        # ----------------------------------------------------

        for column_cells in worksheet.columns:

            max_length = 0

            for cell in column_cells:

                value = (
                    str(cell.value)
                    if cell.value is not None
                    else ""
                )

                max_length = max(
                    max_length,
                    len(value),
                )

            worksheet.column_dimensions[
                column_cells[0].column_letter
            ].width = min(
                max_length + 2,
                45,
            )

    # ========================================================
    # RÉPONSE HTTP
    # ========================================================

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    response["Content-Disposition"] = (
        'attachment; filename="rapport-elisabeth.xlsx"'
    )

    workbook.save(response)

    return response
# ============================================================
# EXPORT PDF
# ============================================================

@api_view(["GET"])
@permission_classes([IsAuthenticated])
def dashboard_pdf(request):

    payload = _dashboard_payload(request)

    summary = payload["summary"]

    response = HttpResponse(
        content_type="application/pdf"
    )

    response[
        "Content-Disposition"
    ] = (
        'attachment; '
        'filename="rapport-elisabeth.pdf"'
    )

    document = SimpleDocTemplate(
        response,
        pagesize=A4,
        rightMargin=25,
        leftMargin=25,
        topMargin=25,
        bottomMargin=25,
    )

    styles = getSampleStyleSheet()

    elements = []

    elements.append(
        Paragraph(
            "La Casa da Festa Elisabeth",
            styles["Title"],
        )
    )

    elements.append(
        Paragraph(
            "Rapport du tableau de bord",
            styles["Heading2"],
        )
    )

    elements.append(
        Spacer(1, 15)
    )

    # ========================================================
    # FILTRES
    # ========================================================

    filter_lines = []

    dashboard_filters = payload.get(
        "filters",
        {}
    )

    if dashboard_filters.get("date"):
        filter_lines.append(
            f"Date : {dashboard_filters['date']}"
        )

    if dashboard_filters.get("jour"):
        filter_lines.append(
            f"Jour : {dashboard_filters['jour']}"
        )

    if dashboard_filters.get("mois"):
        filter_lines.append(
            f"Mois : {dashboard_filters['mois']}"
        )

    if dashboard_filters.get("annee"):
        filter_lines.append(
            f"Année : {dashboard_filters['annee']}"
        )

    if dashboard_filters.get("date_debut"):
        filter_lines.append(
            f"Du : {dashboard_filters['date_debut']}"
        )

    if dashboard_filters.get("date_fin"):
        filter_lines.append(
            f"Au : {dashboard_filters['date_fin']}"
        )

    if filter_lines:

        elements.append(
            Paragraph(
                " | ".join(filter_lines),
                styles["Normal"],
            )
        )

        elements.append(
            Spacer(1, 10)
        )

    # ========================================================
    # SYNTHESE
    # ========================================================

    summary_data = [
        ["Indicateur", "Valeur"],

        [
            "Réservations",
            str(
                summary["total_reservations"]
            ),
        ],

        [
            "Réservations actives",
            str(
                summary["reservations_actives"]
            ),
        ],

        [
            "Réservations annulées",
            str(
                summary["reservations_annulees"]
            ),
        ],

        [
            "Réservations terminées",
            str(
                summary["reservations_terminees"]
            ),
        ],

        [
            "Chiffre d'affaires",
            f'{summary["chiffre_affaires"]:.2f} $',
        ],

        [
            "Total encaissé",
            f'{summary["total_encaisse"]:.2f} $',
        ],

        [
            "Dépenses",
            f'{summary["total_depenses"]:.2f} $',
        ],

        [
            "Remboursements",
            f'{summary["total_rembourse"]:.2f} $',
        ],

        [
            "Reste à recouvrer",
            f'{summary["reste_a_recouvrer"]:.2f} $',
        ],

        [
            "Solde comptes",
            f'{summary["solde_comptes"]:.2f} $',
        ],

        [
            "Résultat net",
            f'{summary["resultat_net"]:.2f} $',
        ],

        [
            "Total entrées",
            f'{summary["total_entrees"]:.2f} $',
        ],

        [
            "Total sorties",
            f'{summary["total_sorties"]:.2f} $',
        ],
    ]

    summary_table = Table(
        summary_data,
        colWidths=[
            100,
            100,
        ],
    )

    summary_table.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.HexColor("#1e293b"),
            ),
            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white,
            ),
            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold",
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.5,
                colors.grey,
            ),
            (
                "PADDING",
                (0, 0),
                (-1, -1),
                6,
            ),
        ])
    )

    elements.append(
        summary_table
    )

    elements.append(
        Spacer(1, 20)
    )

    # ========================================================
    # PAIEMENTS PAR MODE
    # ========================================================

    elements.append(
        Paragraph(
            "Paiements par mode",
            styles["Heading2"],
        )
    )

    payment_rows = [
        [
            "Mode",
            "Montant",
        ]
    ]

    for item in payload[
        "payments_by_method"
    ]:

        payment_rows.append([
            item["label"],
            f'{item["amount"]:.2f} $',
        ])

    payment_table = Table(
        payment_rows,
        colWidths=[
            120,
            100,
        ],
    )

    payment_table.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.HexColor("#1e293b"),
            ),
            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white,
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.5,
                colors.grey,
            ),
            (
                "PADDING",
                (0, 0),
                (-1, -1),
                5,
            ),
        ])
    )

    elements.append(
        payment_table
    )

    elements.append(
        Spacer(1, 20)
    )

    # ========================================================
    # DEPENSES PAR CATEGORIE
    # ========================================================

    elements.append(
        Paragraph(
            "Dépenses par catégorie",
            styles["Heading2"],
        )
    )

    expense_rows = [
        [
            "Catégorie",
            "Montant",
        ]
    ]

    for item in payload[
        "expenses_by_category"
    ]:

        expense_rows.append([
            item["label"],
            f'{item["amount"]:.2f} $',
        ])

    expense_table = Table(
        expense_rows,
        colWidths=[
            120,
            100,
        ],
    )

    expense_table.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.HexColor("#1e293b"),
            ),
            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white,
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.5,
                colors.grey,
            ),
            (
                "PADDING",
                (0, 0),
                (-1, -1),
                5,
            ),
        ])
    )

    elements.append(
        expense_table
    )

    document.build(elements)

    return response



class RefundViewSet(viewsets.ModelViewSet):
    queryset = (
        Refund.objects
        .select_related(
            "payment",
            "reservation",
            "reservation__client",
            "financial_account",
            "created_by",
        )
        .all()
        .order_by("-created_at")
    )

    serializer_class = RefundSerializer
    permission_classes = [IsAuthenticated]

    @transaction.atomic
    def perform_create(self, serializer):
        refund = serializer.save(
            created_by=self.request.user
        )

        return refund

    @action(
        detail=True,
        methods=["post"],
        url_path="valider",
    )
    @transaction.atomic
    def valider(self, request, pk=None):
        refund = (
            Refund.objects
            .select_for_update()
            .select_related(
                "payment",
                "reservation",
                "financial_account",
            )
            .get(pk=pk)
        )

        if refund.status == Refund.Status.VALIDE:
            return Response(
                {
                    "detail": (
                        "Ce remboursement est déjà validé."
                    )
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        if refund.status == Refund.Status.ANNULE:
            return Response(
                {
                    "detail": (
                        "Ce remboursement a été annulé."
                    )
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        self._process_refund_payout(
            refund,
            request.user,
        )

        refund.status = Refund.Status.VALIDE

        refund.save(
            update_fields=[
                "status",
            ]
        )

        if refund.reservation:
            refund.reservation.recalculate_financials()

        self._generate_receipt(refund)

        return Response(
            RefundSerializer(
                refund,
                context={
                    "request": request,
                },
            ).data
        )

    def _process_refund_payout(
        self,
        refund,
        user,
    ):
        payment = (
            Payment.objects
            .select_for_update()
            .get(pk=refund.payment_id)
        )

        account = (
            FinancialAccount.objects
            .select_for_update()
            .get(
                pk=refund.financial_account_id
            )
        )

        if payment.status != Payment.Status.VALIDE:
            raise serializers.ValidationError(
                {
                    "detail": (
                        "Le paiement doit être validé "
                        "avant de pouvoir être remboursé."
                    )
                }
            )

        already_refunded = (
            Refund.objects
            .filter(
                payment=payment,
                status=Refund.Status.VALIDE,
            )
            .exclude(pk=refund.pk)
            .aggregate(
                total=Sum("amount")
            )["total"]
            or Decimal("0.00")
        )

        refundable_amount = (
            payment.amount - already_refunded
        )

        if refund.amount > refundable_amount:
            raise serializers.ValidationError(
                {
                    "amount": (
                        "Le montant du remboursement dépasse "
                        "le montant encore remboursable."
                    )
                }
            )

        if account.balance < refund.amount:
            raise serializers.ValidationError(
                {
                    "detail": (
                        f"Solde insuffisant dans la caisse "
                        f"'{account.name}' pour effectuer "
                        f"ce remboursement. "
                        f"Solde actuel : {account.balance} $."
                    )
                }
            )

        movement_exists = (
            CashMovement.objects.filter(
                refund=refund,
                movement_type=CashMovement.MovementType.SORTIE,
            ).exists()
        )

        if movement_exists:
            return

        account.balance -= refund.amount

        account.save(
            update_fields=[
                "balance",
            ]
        )

        CashMovement.objects.create(
            account=account,
            refund=refund,
            reservation=refund.reservation,
            movement_type=CashMovement.MovementType.SORTIE,
            amount=refund.amount,
            description=(
                f"Sortie Remboursement #{refund.id} "
                f"- Paiement #{payment.id} "
                f"- Réservation "
                f"#{refund.reservation_id}"
            ),
            created_by=user,
        )

    def _generate_receipt(self, refund):
        from .services.pdf_service import (
            generate_refund_receipt_pdf,
        )

        pdf_file = generate_refund_receipt_pdf(
            refund
        )

        refund.receipt_pdf.save(
            pdf_file.name,
            pdf_file,
            save=True,
        )

    @action(
        detail=True,
        methods=["get"],
        url_path="recu",
    )
    def recu(self, request, pk=None):
        refund = (
            Refund.objects
            .select_related(
                "payment",
                "reservation",
                "financial_account",
            )
            .get(pk=pk)
        )

        if not refund.receipt_pdf:
            if refund.status != Refund.Status.VALIDE:
                return Response(
                    {
                        "detail": (
                            "Le reçu PDF sera disponible "
                            "après validation du remboursement."
                        )
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            self._generate_receipt(refund)

            refund.refresh_from_db()

        response = FileResponse(
            refund.receipt_pdf.open("rb"),
            content_type="application/pdf",
        )

        response["Content-Disposition"] = (
            f'inline; filename="'
            f'recu_remboursement_{refund.id}.pdf"'
        )

        return response